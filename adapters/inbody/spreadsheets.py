"""Bounded spreadsheet evidence extraction; only explicit labels and units are mapped."""
import csv
import io
import json
import re
import zipfile
from adapters.inbody.documents import ALIASES, MAX_BYTES, parse_text
from backend.validation import Invalid

VERSION = 'spreadsheet-evidence-v1.0'


def extract_sheet(data, kind):
    if len(data) > MAX_BYTES:
        raise Invalid('Maximum file size is 10 MiB')
    tables = []
    if kind == 'csv':
        try:
            text = data.decode('utf-8-sig')
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except UnicodeError:
            raise Invalid('CSV must use UTF-8 encoding')
        except csv.Error:
            dialect = csv.excel
        try:
            tables = [{'sheet': 'CSV', 'rows': list(csv.reader(io.StringIO(text), dialect))}]
        except csv.Error:
            raise Invalid('Unreadable CSV or cell exceeds size limit')
    elif kind == 'xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(x.file_size for x in archive.infolist()) > 30_000_000:
                    raise Invalid('Workbook decompressed size exceeds limits')
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > 10:
                    raise Invalid('Workbook exceeds 10 sheets')
                for sheet in workbook.worksheets:
                    if sheet.max_row and sheet.max_row > 2000 or sheet.max_column and sheet.max_column > 200:
                        raise Invalid('Sheet exceeds row/column limits')
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        rows.append([str(value) if value is not None else '' for value in row])
                        if len(rows) > 2000 or len(row) > 200:
                            raise Invalid('Sheet exceeds row/column limits')
                    tables.append({'sheet': sheet.title, 'rows': rows})
            finally:
                workbook.close()
        except Invalid:
            raise
        except Exception:
            raise Invalid('Unreadable XLSX workbook')
    elif kind == 'xls':
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if workbook.nsheets > 10:
                    raise Invalid('Workbook exceeds 10 sheets')
                for sheet in workbook.sheets():
                    if sheet.nrows > 2000 or sheet.ncols > 200:
                        raise Invalid('Sheet exceeds row/column limits')
                    tables.append({'sheet': sheet.name, 'rows': [[str(v) for v in sheet.row_values(i)] for i in range(sheet.nrows)]})
            finally:
                workbook.release_resources()
        except Invalid:
            raise
        except Exception:
            raise Invalid('Unreadable XLS workbook')
    else:
        raise Invalid('Unsupported spreadsheet format')
    raw = json.dumps(tables, ensure_ascii=False)
    if len(raw) > 200000 or sum(len(t['rows']) for t in tables) > 2000 or sum(len(row) for t in tables for row in t['rows']) > 40000:
        raise Invalid('Spreadsheet evidence exceeds limits')
    lines, warnings, identities = [], [], set()
    for table in tables:
        rows = [row for row in table['rows'] if any(str(v).strip() for v in row)]
        if not rows:
            continue
        header = [str(v).strip().lower() for v in rows[0]]
        for key in ('patient_id','patient id'):
            if key in header:
                column = header.index(key)
                identities.update(str(row[column]).strip() for row in rows[1:] if len(row)>column and str(row[column]).strip())
        if all(field in header for field in ('metric', 'value', 'unit')):
            for row in rows[1:]:
                if len(row) <= max(header.index(field) for field in ('metric', 'value', 'unit')):
                    continue
                metric, value, unit = [str(row[header.index(field)]).strip() for field in ('metric', 'value', 'unit')]
                if metric in ALIASES:
                    lines.append(ALIASES[metric][0] + ': ' + value + ' ' + unit)
        elif len(rows) == 2:
            # A single exported test row; headers must explicitly include units.
            for label, value in zip(rows[0], rows[1]):
                lines.append(str(label).strip() + ' ' + str(value).strip())
        else:
            warnings.append('Unsupported or multi-test table retained in full; transcribe one test with explicit units and identity.')
    if len(identities)>1:
        raise Invalid('Export contains multiple patient identifiers; export one patient/test for import.')
    if identities:
        identity = next(iter(identities))
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,80}',identity):
            raise Invalid('Export patient ID cannot be matched safely; use a selected export with a text identifier.')
        lines.append('Patient ID: '+identity)
    parsed = parse_text('\n'.join(lines), allow_empty=True)
    parsed.update({'parser_version': VERSION, 'raw_text': raw, 'tables': tables, 'pages': [], 'barcodes': [],
                   'warnings': warnings, 'coverage': 'Explicit metric/value/unit rows or a single test row with labeled units. Vendor-specific mapping is not inferred.'})
    return parsed
